"""Fuentes del reporte: catalogo de modelos .rvt de VENTAS en ACC (Data Management, 2-legged, gratis)
y Excel de integrantes. Copiado de PUBLICACIONES_VENTAS (pub_sync/fuentes.py) para usar exactamente
las mismas reglas de nombres, disciplinas y personas.
"""
import csv
import json
import logging
import re
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path

log = logging.getLogger("audit_sync.fuentes")


# ================================================================
# NORMALIZACION (misma logica que las consultas de Power Query)
# ================================================================

def modelo_key(nombre):
    """'11_ARQ_CONJUNTO.rvt' -> '11_ARQ_CONJUNTO'  (NormalizarModelo)"""
    s = re.sub(r"[\x00-\x1f\x7f]", "", str(nombre or "")).strip().upper()
    return s[:-4] if s.endswith(".RVT") else s


def modelo_sin_ext(nombre):
    s = str(nombre or "").strip()
    return s[:-4] if s.lower().endswith(".rvt") else s


def proyecto_norm(v):
    """'GRANJAS_JESSY' -> 'GRANJAS JESSY'"""
    return re.sub(r"\s+", " ", str(v or "").replace("_", " ")).strip()


def persona_norm(v):
    """NormalizarPersona: mayusculas, sin acentos (incluye N), puntuacion como espacio."""
    s = re.sub(r"[\x00-\x1f\x7f]", "", str(v or "")).strip().upper()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = re.sub(r"[.,;:\-_/\\]", " ", s)
    return " ".join(s.split())


def persona_compacta(v):
    """CompactarPersona: 'Juan Pablo Reyes' y 'JUANPABLOREYES' dan la misma llave."""
    return persona_norm(v).replace(" ", "")


def disciplina(modelo):
    """Misma regla que HISTORIAL_PUBLICACIONES_MODELOS (por palabras en el nombre del archivo)."""
    a = str(modelo or "").upper()
    if "ARQ" in a or "ARC" in a:
        return "ARQUITECTURA"
    if "EST" in a or "STR" in a:
        return "ESTRUCTURA"
    if "ELE" in a:
        return "ELECTRICA"
    if "ESP" in a or "SPE" in a:
        return "ESPECIALES"
    if "MEC" in a:
        return "MECANICA"
    if "PLO" in a or "PLU" in a:
        return "PLOMERIA"
    return "SIN DISCIPLINA"


def parse_fecha(texto, tz):
    """ISO 8601 con o sin zona ('2026-05-04T21:34:07.000Z', '...-06:00', 7 decimales)
    -> datetime en la zona local del reporte (tz). Sin zona se asume UTC."""
    if not texto:
        return None
    s = str(texto).strip().replace("Z", "+00:00")
    m = re.match(r"^(\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2})?)(\.\d+)?([+-]\d{2}:?\d{2})?$", s)
    if not m:
        return None
    base, frac, zona = m.groups()
    frac = (frac or ".0")[:7].ljust(7, "0")
    dt = datetime.fromisoformat(base.replace(" ", "T") + frac + (zona or "+00:00"))
    return dt.astimezone(tz).replace(tzinfo=None)


# ================================================================
# 1. ACC: CATALOGO DE MODELOS
# ================================================================

class Cache:
    def __init__(self, ruta):
        self.p = Path(ruta)
        try:
            self.d = json.loads(self.p.read_text(encoding="utf-8"))
        except Exception:
            self.d = {}

    def save(self):
        self.p.parent.mkdir(parents=True, exist_ok=True)
        self.p.write_text(json.dumps(self.d, ensure_ascii=False), encoding="utf-8")


def _buscar_subcarpetas(aps, pid, folder_id, nombre, ruta, max_prof=6):
    """Todas las carpetas llamadas 'nombre' (p.ej. 011_WIP) dentro del proyecto, en cualquier nivel
    (p.ej. HOWA/01_D&I/011_WIP). No se detiene en la primera: un proyecto puede tener mas de una.
    No se desciende dentro de una 011_WIP encontrada (su contenido se recorre despues)."""
    nivel, obj, encontradas = [(folder_id, ruta)], nombre.strip().lower(), []
    for _ in range(max_prof):
        sig = []
        for fid, r in nivel:
            carpetas, _ = aps.contenido(pid, fid)
            for c in carpetas:
                rc = f"{r}/{c['name']}"
                if c["name"].strip().lower() == obj:
                    encontradas.append({**c, "ruta": rc})
                else:
                    sig.append((c["id"], rc))
        nivel = sig
        if not nivel:
            break
    return encontradas


def _rvt_recursivo(aps, pid, folder_id, excluir_re, ruta):
    carpetas, items = aps.contenido(pid, folder_id)
    for it in items:
        if it["name"].lower().endswith(".rvt") and not (excluir_re and re.search(excluir_re, it["name"], re.I)):
            yield {**it, "ruta": ruta}
    for c in carpetas:
        yield from _rvt_recursivo(aps, pid, c["id"], excluir_re, f"{ruta}/{c['name']}")


def catalogo_modelos(aps, cfg, cache_path):
    """-> ({item_id: {proyecto, modelo, ruta, archivo}}, avisos)
    Todos los .rvt dentro de 011_WIP de cada proyecto de VENTAS (misma busqueda que PUBLICACIONES,
    sin pedir versiones). Si ACC no responde, se usa el ultimo catalogo guardado en cache."""
    cache = Cache(cache_path)
    try:
        a = cfg["aps"]
        hub, pid = aps.con_b(a["account_id"]), aps.con_b(a["project_id"])
        raices = [n.lower() for n in cfg.get("raiz_nombres", ["Project Files"])]
        tops = aps.top_folders(hub, pid)
        raiz = next((t for t in tops if t["name"].strip().lower() in raices), None)
        if not raiz:
            raise ValueError(f"No se encontro la carpeta raiz {raices} entre {[t['name'] for t in tops]}")
        sub = cfg.get("subcarpeta_datos", "011_WIP")
        excl = {x.lower() for x in cfg.get("excluir_proyectos", [])}
        excluir_re = cfg.get("excluir_rvt_regex")
        modelos = {}
        proyectos, _ = aps.contenido(pid, raiz["id"])
        for p in proyectos:
            nombre_p = p["name"].strip()
            if nombre_p.upper().startswith("Z_") or nombre_p.lower() in excl:
                continue
            ruta_p = f"{raiz['name']}/{nombre_p}"
            wips = _buscar_subcarpetas(aps, pid, p["id"], sub, ruta_p) if sub else [{**p, "ruta": ruta_p}]
            for it in (it for w in wips for it in _rvt_recursivo(aps, pid, w["id"], excluir_re, w["ruta"])):
                modelos[it["item_id"]] = {"proyecto": nombre_p, "modelo": modelo_sin_ext(it["name"]),
                                          "ruta": it["ruta"], "archivo": it["name"]}
        log.info("ACC: %d modelos .rvt en %s de %d proyectos", len(modelos), sub, len(proyectos))
        cache.d = {"guardado": datetime.now().isoformat(timespec="seconds"), "modelos": modelos}
        cache.save()
        return modelos, []
    except Exception as e:
        if cache.d.get("modelos"):
            log.warning("No se pudo leer el catalogo de ACC (%s). Se usa el del %s.", e, cache.d.get("guardado"))
            return cache.d["modelos"], [f"ACC no respondio ({str(e)[:120]}); se uso el catalogo de modelos del {cache.d.get('guardado')}."]
        raise


# ================================================================
# 2. EQUIPOS E INTEGRANTES
# ================================================================

def leer_equipos(ruta, hoja="Integrantes"):
    """-> lista de {integrante, equipo, compacta} (distinta por llave compacta)."""
    from openpyxl import load_workbook
    ws = load_workbook(ruta, data_only=True, read_only=True)[hoja]
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c or "").strip() for c in filas[0]]
    ie, ii = enc.index("Equipo"), enc.index("Integrante")
    out, vistos = [], set()
    for r in filas[1:]:
        eq = str(r[ie] or "").replace("\xa0", " ").strip()
        it = str(r[ii] or "").replace("\xa0", " ").strip()
        if not eq or not it:
            continue
        k = persona_compacta(it)
        if k in vistos:
            continue
        vistos.add(k)
        out.append({"integrante": it, "equipo": eq, "compacta": k})
    return out


def _variantes_login(integrante):
    """Formas en que un nombre del Excel puede aparecer como usuario de Autodesk:
    cualquier combinacion en orden de 2 o mas de sus palabras, pegadas.
    'Mario Alberto Sanchez Munoz' -> MARIOSANCHEZ, MARIOALBERTOSANCHEZ, MARIOSANCHEZMUNOZ, ..."""
    from itertools import combinations
    palabras = persona_norm(integrante).split()
    out = set()
    for n in range(2, len(palabras) + 1):
        for combo in combinations(palabras, n):
            out.add("".join(combo))
    return out


class Personas:
    """Resuelve un nombre (usuario Revit o de ACC) contra el listado oficial del Excel.

    Prioridad:
      1. Alias de config.json (para casos que no se resuelven solos).
      2. Nombre completo igual al del Excel (sin acentos, mayusculas ni espacios).
      3. Usuario de Autodesk tipo 'mariosanchezgcp': se quita el sufijo 'gcp' y se busca
         el integrante cuyo nombre forme ese usuario (MARIO + SANCHEZ). Solo se acepta si
         coincide con UNA sola persona del listado; si hay duplicados no se adivina.
    Sin coincidencia: nombre original y SIN EQUIPO (queda registrado en no_resueltos)."""

    SUFIJOS_LOGIN = ("GCP",)

    def __init__(self, equipos, alias=None):
        self.por_k = {e["compacta"]: e for e in equipos}
        self.alias = {persona_compacta(k): v for k, v in (alias or {}).items()}
        variantes = {}
        for e in equipos:
            for v in _variantes_login(e["integrante"]):
                variantes.setdefault(v, []).append(e)
        self.por_login = {v: es[0] for v, es in variantes.items() if len(es) == 1}
        self.no_resueltos = set()

    def _por_login(self, k):
        for suf in self.SUFIJOS_LOGIN:
            if k.endswith(suf) and len(k) > len(suf):
                e = self.por_login.get(k[: -len(suf)])
                if e:
                    return e
        return self.por_login.get(k)

    def resolver(self, nombre):
        k = persona_compacta(nombre)
        destino = self.alias.get(k)
        if destino:
            k = persona_compacta(destino)
        e = self.por_k.get(k) or (self._por_login(k) if k else None)
        if e:
            return e["integrante"], e["equipo"]
        if nombre:
            self.no_resueltos.add(nombre)
        return (nombre or None), "SIN EQUIPO"
